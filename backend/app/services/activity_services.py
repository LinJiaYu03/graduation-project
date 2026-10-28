import base64
import json
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from functools import cmp_to_key
from io import BytesIO
from urllib.parse import quote_plus

import qrcode
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.core.rbac import Perm, assert_permission, has_permission
from app.core.respones import page_result, success_response
from app.models.activity import Activity, ActivityCheckin, ActivityRegistration
from app.models.club import Club, ClubMember
from app.models.user import User
from app.schemas.activity import (
    ActivityCommentRequest,
    ActivityCreateRequest,
    ActivityReviewRequest,
    ActivitySearchRequest,
    ActivityUpdateRequest,
    RegistrationReviewRequest,
    _number,
    _split_urls,
    serialize_activity,
    serialize_activity_dto,
    serialize_registration,
    serialize_user_dto,
)
from app.utils.time_util import dt_to_millis


class ExportError(Exception):
    """导出接口错误（对齐 Java 直接写 HttpServletResponse 状态码的行为）"""

    def __init__(self, status_code: int, message: str):
        super().__init__(message)
        self.status_code = status_code
        self.message = message


def _has_text(v) -> bool:
    """对齐 Spring StringUtils.hasText"""
    return v is not None and str(v).strip() != ""


def _to_long(v) -> int | None:
    if v is None:
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _round_half_up(value, places: int) -> Decimal:
    quant = Decimal(1).scaleb(-places)
    return Decimal(str(value)).quantize(quant, rounding=ROUND_HALF_UP)


def _approved_count(db: Session, activity_id) -> int:
    return (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.status == "已通过",
        )
        .count()
    )


def refresh_activity_status_by_time(db: Session, activity: Activity) -> None:
    """对齐 Java refreshActivityStatusByTime，会根据当前时间刷新状态并持久化。"""
    if activity is None:
        return
    if activity.audit_status is None or activity.audit_status != 1:
        return

    now = datetime.now()
    reg_start = activity.registration_start_time
    reg_end = activity.registration_end_time
    start = activity.start_time
    end = activity.end_time

    if reg_start is None or reg_end is None or start is None or end is None:
        db.commit()
        return

    if now > reg_start and now < reg_end:
        new_status = "报名中"
    elif (now > reg_end and now < start) or now < reg_start:
        new_status = "等待中"
    elif now > start and now < end:
        new_status = "进行中"
    elif now > end:
        new_status = "已结束"
    else:
        new_status = activity.status

    if new_status is not None and new_status != activity.status:
        activity.status = new_status
    db.commit()


def _status_priority(status) -> int:
    if status == "报名中":
        return 1
    if status == "进行中":
        return 2
    if status == "已结束":
        return 3
    return 4


# ---------------------------------------------------------------------------
# 创建 / 编辑 / 审核
# ---------------------------------------------------------------------------

def create_activity(
    db: Session, request: ActivityCreateRequest, current_user: User
):
    if not _has_text(request.name):
        raise ValueError("活动名称不能为空")
    if request.clubId is None:
        raise ValueError("主办方社团ID不能为空")

    club = db.query(Club).filter(Club.id == request.clubId).first()
    if club is None:
        raise ValueError("主办方社团不存在")

    assert_permission(
        db,
        current_user,
        Perm.ACTIVITY_CREATE,
        club_id=club.id,
        message="无权限在该社团下创建活动",
    )

    activity = Activity()
    activity.name = request.name
    activity.description = request.description
    activity.activity_type = request.activityType
    activity.location = request.location
    activity.club_id = request.clubId
    activity.notice = request.notice
    activity.registration_start_time = request.registrationStartTime
    activity.registration_end_time = request.registrationEndTime
    activity.start_time = request.startTime
    activity.end_time = request.endTime
    activity.max_participants = request.maxParticipants
    activity.need_audit = request.needAudit if request.needAudit is not None else False

    if request.imageUrls:
        activity.image_urls = ",".join(request.imageUrls)

    # 对齐 ActivityServiceImpl.createActivity
    activity.activity_level = club.level_tag
    activity.status = "审核中"
    activity.audit_status = 0
    activity.score = "0"
    activity.is_published = False

    db.add(activity)
    db.commit()
    db.refresh(activity)
    return success_response(
        data={"activity": serialize_activity(activity)}, message="创建活动成功"
    )


def update_activity(
    db: Session, request: ActivityUpdateRequest, current_user: User
):
    if request.activityId is None:
        raise ValueError("活动ID不能为空")

    activity = db.query(Activity).filter(Activity.id == request.activityId).first()
    if activity is None:
        raise ValueError("活动不存在")

    assert_permission(
        db,
        current_user,
        Perm.ACTIVITY_UPDATE,
        club_id=activity.club_id,
        message="无权限编辑该活动",
    )

    if _has_text(request.name):
        activity.name = request.name
    if request.description is not None:
        activity.description = request.description
    if request.activityType is not None:
        activity.activity_type = request.activityType
    if request.location is not None:
        activity.location = request.location
    if request.notice is not None:
        activity.notice = request.notice
    if request.registrationStartTime is not None:
        activity.registration_start_time = request.registrationStartTime
    if request.registrationEndTime is not None:
        activity.registration_end_time = request.registrationEndTime
    if request.startTime is not None:
        activity.start_time = request.startTime
    if request.endTime is not None:
        activity.end_time = request.endTime
    if request.maxParticipants is not None:
        activity.max_participants = request.maxParticipants
    if request.needAudit is not None:
        activity.need_audit = request.needAudit
    if request.imageUrls is not None:
        if request.imageUrls:
            activity.image_urls = ",".join(request.imageUrls)
        else:
            activity.image_urls = None

    # 编辑后重新进入审核状态
    activity.audit_status = 0
    activity.status = "审核中"
    activity.audit_user_id = None
    activity.audit_time = None

    db.commit()
    db.refresh(activity)
    return success_response(
        data={"activity": serialize_activity(activity)},
        message="编辑活动成功，活动已重新进入审核状态",
    )


def pending_activities(
    db: Session, current_user: User, current_page: int, page_size: int
):
    assert_permission(db, current_user, Perm.ACTIVITY_REVIEW)

    query = (
        db.query(Activity)
        .filter(Activity.audit_status == 0)
        .order_by(Activity.create_time.desc())
    )
    total = query.count()
    records = query.offset((current_page - 1) * page_size).limit(page_size).all()
    return success_response(
        data={
            "page": page_result(
                total,
                current_page,
                page_size,
                [serialize_activity_dto(a) for a in records],
            )
        },
        message="查询成功",
    )


def review_activity(
    db: Session, request: ActivityReviewRequest, current_user: User
):
    if request.activityId is None or request.pass_ is None:
        raise ValueError("参数不完整")

    assert_permission(db, current_user, Perm.ACTIVITY_REVIEW)

    activity = db.query(Activity).filter(Activity.id == request.activityId).first()
    if activity is None:
        raise ValueError("活动不存在")

    if request.pass_ is True:
        activity.audit_status = 1
    else:
        activity.audit_status = 2
        activity.status = "已拒绝"
    activity.audit_user_id = current_user.id
    activity.audit_time = datetime.now()

    if activity.audit_status == 1:
        refresh_activity_status_by_time(db, activity)
    else:
        db.commit()
    return success_response(message="审核成功")


# ---------------------------------------------------------------------------
# 报名审核
# ---------------------------------------------------------------------------

def pending_registrations(db: Session, activity_id, current_user: User):
    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ValueError("活动不存在")

    assert_permission(
        db,
        current_user,
        Perm.REGISTRATION_VIEW,
        club_id=activity.club_id,
        message="无权限查看该活动报名信息",
    )

    records = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.status == "待审核",
        )
        .order_by(ActivityRegistration.registration_time.asc())
        .all()
    )
    return [serialize_registration(r) for r in records]


def review_registration(
    db: Session, request: RegistrationReviewRequest, current_user: User
):
    if request.activityId is None or request.userId is None or request.pass_ is None:
        raise ValueError("参数不完整")

    activity = db.query(Activity).filter(Activity.id == request.activityId).first()
    if activity is None:
        raise ValueError("活动不存在")

    assert_permission(
        db,
        current_user,
        Perm.REGISTRATION_REVIEW,
        club_id=activity.club_id,
        message="无审核权限",
    )

    registration = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == request.activityId,
            ActivityRegistration.user_id == request.userId,
        )
        .first()
    )
    if registration is None:
        raise ValueError("报名记录不存在")

    old_status = registration.status
    if request.pass_ is True:
        registration.status = "已通过"
    else:
        registration.status = "已拒绝"
    registration.audit_user_id = current_user.id
    registration.audit_time = datetime.now()
    db.commit()

    if request.pass_ is True and old_status != "已通过":
        activity.current_participants = _approved_count(db, request.activityId)
        db.commit()
    elif request.pass_ is not True and old_status == "已通过":
        activity.current_participants = _approved_count(db, request.activityId)
        db.commit()
    return success_response(message="审核成功")


# ---------------------------------------------------------------------------
# 查询
# ---------------------------------------------------------------------------

def search_activities(db: Session, request: ActivitySearchRequest):
    current = request.currentPage if (request.currentPage and request.currentPage > 0) else 1
    size = request.pageSize if (request.pageSize and request.pageSize > 0) else 10

    query = db.query(Activity)
    # 排除已拒绝的活动
    query = query.filter(Activity.audit_status != 2)
    query = query.filter(Activity.is_published.is_(True))

    if _has_text(request.keyword):
        kw = f"%{request.keyword}%"
        query = query.filter(
            or_(
                Activity.name.like(kw),
                Activity.description.like(kw),
                Activity.location.like(kw),
            )
        )
    if _has_text(request.status):
        query = query.filter(Activity.status == request.status)
    if _has_text(request.category):
        query = query.filter(Activity.activity_type == request.category)
    if _has_text(request.level):
        query = query.filter(Activity.activity_level == request.level)

    query = query.order_by(Activity.create_time.desc())

    total = query.count()
    records = query.offset((current - 1) * size).limit(size).all()
    for activity in records:
        refresh_activity_status_by_time(db, activity)
    return success_response(
        data={
            "page": page_result(
                total, current, size, [serialize_activity_dto(a) for a in records]
            )
        },
        message="查询成功",
    )


def hot_activities(db: Session):
    activities = (
        db.query(Activity)
        .filter(Activity.audit_status == 1, Activity.is_published.is_(True))
        .all()
    )

    for activity in activities:
        activity.current_participants = _approved_count(db, activity.id)
        refresh_activity_status_by_time(db, activity)

    hot = [
        a
        for a in activities
        if a.status in ("报名中", "进行中", "已结束")
    ]

    def _compare(a1, a2):
        p1 = _status_priority(a1.status)
        p2 = _status_priority(a2.status)
        if p1 != p2:
            return -1 if p1 < p2 else 1
        c1 = a1.current_participants if a1.current_participants is not None else 0
        c2 = a2.current_participants if a2.current_participants is not None else 0
        if c1 == c2:
            return 0
        return -1 if c2 < c1 else 1

    hot.sort(key=cmp_to_key(_compare))
    return success_response(
        data={"activities": [serialize_activity_dto(a) for a in hot[:6]]},
        message="查询成功",
    )


def activities_by_club(db: Session, club_id, current_page: int, page_size: int):
    club = db.query(Club).filter(Club.id == club_id).first()
    if club is None:
        raise ValueError("社团不存在")

    query = (
        db.query(Activity)
        .filter(Activity.club_id == club_id, Activity.is_published.is_(True))
        .order_by(Activity.create_time.desc())
    )
    total = query.count()
    records = query.offset((current_page - 1) * page_size).limit(page_size).all()
    for activity in records:
        refresh_activity_status_by_time(db, activity)
        activity.current_participants = _approved_count(db, activity.id)
    return success_response(
        data={
            "page": page_result(
                total,
                current_page,
                page_size,
                [serialize_activity_dto(a) for a in records],
            )
        },
        message="查询成功",
    )


def get_activity_detail(db: Session, activity_id):
    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ValueError("活动不存在")

    refresh_activity_status_by_time(db, activity)
    activity.current_participants = _approved_count(db, activity_id)

    activity_map = {
        "id": activity.id,
        "name": activity.name,
        "clubId": activity.club_id,
        "description": activity.description,
        "location": activity.location,
        "notice": activity.notice,
        "registrationStartTime": dt_to_millis(activity.registration_start_time),
        "registrationEndTime": dt_to_millis(activity.registration_end_time),
        "startTime": dt_to_millis(activity.start_time),
        "endTime": dt_to_millis(activity.end_time),
        "maxParticipants": activity.max_participants,
        "currentParticipants": activity.current_participants,
        "needAudit": activity.need_audit,
        "activityType": activity.activity_type,
        "activityLevel": activity.activity_level,
        "status": activity.status,
        "score": _number(activity.score),
        "auditStatus": activity.audit_status,
        "auditTime": activity.audit_time,
        "auditUserId": activity.audit_user_id,
        "isPublished": activity.is_published,
        "createTime": activity.create_time,
        "updateTime": activity.update_time,
    }

    commented_registrations = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.status == "已通过",
            ActivityRegistration.comment.isnot(None),
            ActivityRegistration.comment != "",
        )
        .all()
    )

    commented_members = []
    for reg in commented_registrations:
        user = db.query(User).filter(User.id == reg.user_id).first()
        if user is None:
            continue
        commented_members.append(
            {
                "userId": user.id,
                "realName": user.real_name,
                "studentNo": user.student_no,
                "avatarUrl": user.avatar_url,
                "comment": reg.comment,
                "score": _number(reg.score),
                "commentTime": dt_to_millis(reg.update_time),
            }
        )
    activity_map["commentedMembers"] = commented_members

    data = {"activity": activity_map}
    if _has_text(activity.image_urls):
        data["imageUrls"] = _split_urls(activity.image_urls)
    return success_response(data=data, message="查询成功")


# ---------------------------------------------------------------------------
# 报名 / 评价
# ---------------------------------------------------------------------------

def register_activity(
    db: Session, activity_id, request, current_user: User
):
    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ValueError("活动不存在")
    if activity.audit_status is None or activity.audit_status != 1:
        raise ValueError("活动尚未通过审核，无法报名")
    if activity.is_published is not True:
        raise ValueError("活动已下架，无法报名")

    now = datetime.now()
    if activity.registration_start_time is None or activity.registration_end_time is None:
        raise ValueError("活动报名时间未配置，无法报名")
    if now < activity.registration_start_time or now > activity.registration_end_time:
        raise ValueError("当前不在报名时间范围内")

    existing = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.user_id == current_user.id,
        )
        .first()
    )
    if existing is not None:
        raise ValueError("您已报名该活动")

    if activity.max_participants is not None and activity.max_participants > 0:
        if _approved_count(db, activity_id) >= activity.max_participants:
            raise ValueError("报名人数已达上限")

    registration = ActivityRegistration()
    registration.activity_id = activity_id
    registration.user_id = current_user.id
    if activity.need_audit is True:
        registration.status = "待审核"
    else:
        registration.status = "已通过"
    registration.registration_time = now
    db.add(registration)
    db.commit()
    db.refresh(registration)

    if activity.need_audit is not True:
        activity.current_participants = _approved_count(db, activity_id)
        db.commit()

    return success_response(
        data={"registration": serialize_registration(registration)}, message="报名成功"
    )


def comment_activity(
    db: Session, request: ActivityCommentRequest, current_user: User
):
    if request.activityId is None:
        raise ValueError("参数不完整")
    if request.score is None and (request.comment is None or request.comment.strip() == ""):
        raise ValueError("请至少提供评分或文字评价")

    activity = db.query(Activity).filter(Activity.id == request.activityId).first()
    if activity is None:
        raise ValueError("活动不存在")

    if activity.end_time is None:
        raise ValueError("活动结束时间未配置，无法评分")
    now = datetime.now()
    if not now > activity.end_time:
        raise ValueError("活动结束后才能进行评分")

    registration = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == request.activityId,
            ActivityRegistration.user_id == current_user.id,
        )
        .first()
    )
    if registration is None:
        raise ValueError("未找到报名记录，不能评价")
    if registration.status != "已通过":
        raise ValueError("报名未通过，不能评价")

    checkin = (
        db.query(ActivityCheckin)
        .filter(
            ActivityCheckin.activity_id == request.activityId,
            ActivityCheckin.user_id == current_user.id,
        )
        .first()
    )
    if checkin is None:
        raise ValueError("您未签到该活动，不能评价")

    if registration.score is not None or (
        registration.comment is not None and registration.comment.strip() != ""
    ):
        raise ValueError("您已经评价过该活动，不能重复评价")

    if request.score is not None:
        score = Decimal(request.score)
        if score < Decimal(0) or score > Decimal(5):
            raise ValueError("评分范围应在0-5之间")
        registration.score = score

    if request.comment is not None and request.comment.strip() != "":
        comment = request.comment.strip()
        if len(comment) > 1000:
            raise ValueError("文字评价不能超过1000字符")
        registration.comment = comment

    db.commit()

    scored_list = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == request.activityId,
            ActivityRegistration.status == "已通过",
            ActivityRegistration.score.isnot(None),
        )
        .all()
    )
    if scored_list:
        total_score = sum(
            (Decimal(r.score) for r in scored_list if r.score is not None), Decimal(0)
        )
        average = (total_score / Decimal(len(scored_list))).quantize(
            Decimal("0.1"), rounding=ROUND_HALF_UP
        )
        activity.score = str(average)
    else:
        activity.score = "0"
    db.commit()
    return success_response(message="评价成功")


# ---------------------------------------------------------------------------
# 报名状态 / 统计
# ---------------------------------------------------------------------------

def registration_status(db: Session, activity_id, current_user: User):
    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ValueError("活动不存在")

    assert_permission(
        db,
        current_user,
        Perm.REGISTRATION_VIEW,
        club_id=activity.club_id,
        message="无权限查看该活动报名信息",
    )

    registrations = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.status == "待审核",
        )
        .order_by(ActivityRegistration.registration_time.asc())
        .all()
    )
    if not registrations:
        return success_response(data={"items": []}, message="查询成功")

    user_ids = list({r.user_id for r in registrations})
    users = db.query(User).filter(User.id.in_(user_ids)).all()
    return success_response(
        data={"items": [serialize_user_dto(u) for u in users]}, message="查询成功"
    )


def check_registration(db: Session, activity_id, user_id):
    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ValueError("活动不存在")

    registration = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.user_id == user_id,
        )
        .first()
    )

    result = {}
    if registration is None:
        result["isRegistered"] = False
        result["status"] = None
    else:
        status = registration.status
        if status == "已通过":
            checkin = (
                db.query(ActivityCheckin)
                .filter(
                    ActivityCheckin.activity_id == activity_id,
                    ActivityCheckin.user_id == user_id,
                )
                .first()
            )
            if checkin is not None:
                status = "签到成功"

        result["isRegistered"] = True
        result["status"] = status
        result["registrationId"] = registration.id
        result["registrationTime"] = dt_to_millis(registration.registration_time)
        result["auditTime"] = dt_to_millis(registration.audit_time)
        result["score"] = _number(registration.score)
        result["isAccepted"] = (
            registration.status == "已通过" or status == "签到成功"
        )

    return success_response(data={"registration": result}, message="查询成功")


def _activity_statistics_base(db: Session, activity_id, current_user: User):
    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ValueError("活动不存在")
    assert_permission(
        db,
        current_user,
        Perm.ACTIVITY_STATS,
        club_id=activity.club_id,
        message="无权限查看该活动统计信息",
    )
    return activity


def activity_statistics(db: Session, activity_id, current_user: User):
    activity = _activity_statistics_base(db, activity_id, current_user)

    pending = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.status == "待审核",
        )
        .count()
    )
    approved = _approved_count(db, activity_id)
    rejected = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.status == "已拒绝",
        )
        .count()
    )

    status_stats = {
        "pending": pending,
        "approved": approved,
        "rejected": rejected,
        "total": pending + approved + rejected,
    }

    success_rate = 0.0
    total = pending + approved + rejected
    if total > 0:
        success_rate = float(_round_half_up(approved / total * 100, 2))

    statistics = {
        "activityInfo": serialize_activity_dto(activity, image_list=False),
        "statusStats": status_stats,
        "successRate": success_rate,
        "maxParticipants": activity.max_participants,
        "currentParticipants": activity.current_participants,
    }
    return success_response(data={"statistics": statistics}, message="查询成功")


def checkin_statistics(db: Session, activity_id, current_user: User):
    activity = _activity_statistics_base(db, activity_id, current_user)

    approved = _approved_count(db, activity_id)
    checked_in = (
        db.query(ActivityCheckin)
        .filter(ActivityCheckin.activity_id == activity_id)
        .count()
    )
    not_checked_in = approved - checked_in

    checkin_rate = 0.0
    if approved > 0:
        checkin_rate = float(_round_half_up(checked_in / approved * 100, 2))

    statistics = {
        "activityInfo": serialize_activity_dto(activity, image_list=False),
        "checkedIn": checked_in,
        "notCheckedIn": not_checked_in,
        "totalApproved": approved,
        "checkinRate": checkin_rate,
    }
    return success_response(data={"statistics": statistics}, message="查询成功")


def score_statistics(db: Session, activity_id, current_user: User):
    activity = _activity_statistics_base(db, activity_id, current_user)

    scored_registrations = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.status == "已通过",
            ActivityRegistration.score.isnot(None),
        )
        .all()
    )

    score_05_10 = 0
    score_15_20 = 0
    score_25_30 = 0
    score_35_40 = 0
    score_45_50 = 0

    total_score = Decimal(0)
    scored_count = len(scored_registrations)

    for reg in scored_registrations:
        score = reg.score
        if score is not None:
            total_score += Decimal(score)
            value = float(score)
            if 0.5 <= value <= 1.0:
                score_05_10 += 1
            elif 1.5 <= value <= 2.0:
                score_15_20 += 1
            elif 2.5 <= value <= 3.0:
                score_25_30 += 1
            elif 3.5 <= value <= 4.0:
                score_35_40 += 1
            elif 4.5 <= value <= 5.0:
                score_45_50 += 1

    average_score = Decimal(0)
    if scored_count > 0:
        average_score = (total_score / Decimal(scored_count)).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        )

    total_approved = _approved_count(db, activity_id)
    not_scored_count = total_approved - scored_count

    score_rate = 0.0
    if total_approved > 0:
        score_rate = float(_round_half_up(scored_count / total_approved * 100, 2))

    statistics = {
        "activityInfo": serialize_activity_dto(activity, image_list=False),
        "scoreDistribution": {
            "0.5-1.0": score_05_10,
            "1.5-2.0": score_15_20,
            "2.5-3.0": score_25_30,
            "3.5-4.0": score_35_40,
            "4.5-5.0": score_45_50,
        },
        "averageScore": float(average_score),
        "scoredCount": scored_count,
        "notScoredCount": not_scored_count,
        "totalApproved": total_approved,
        "scoreRate": score_rate,
    }
    return success_response(data={"statistics": statistics}, message="查询成功")


# ---------------------------------------------------------------------------
# Excel 导出
# ---------------------------------------------------------------------------

def _gender_to_string(gender) -> str:
    if gender is None:
        return ""
    return "女" if gender == 0 else "男"

def _new_workbook(sheet_name: str, headers: list[str]):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_font = Font(bold=True)
    header_fill = PatternFill(
        fill_type="solid", fgColor="D9D9D9", start_color="D9D9D9", end_color="D9D9D9"
    )
    thin = Side(style="thin")
    border = Border(bottom=thin, top=thin, left=thin, right=thin)
    center = Alignment(horizontal="center")

    for index, header in enumerate(headers):
        cell = ws.cell(row=1, column=index + 1, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    for index in range(len(headers)):
        ws.column_dimensions[
            ws.cell(row=1, column=index + 1).column_letter
        ].width = 4000 / 256

    return wb, ws


def _workbook_bytes(wb) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _export_context(db: Session, activity_id, current_user: User):
    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ExportError(404, "活动不存在")
    if not has_permission(
        db, current_user.id, Perm.ACTIVITY_EXPORT, activity.club_id
    ):
        raise ExportError(403, "无权限导出")
    return activity


def _approved_registrations(db: Session, activity_id):
    return (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.status == "已通过",
        )
        .all()
    )


def _fill_user_columns(ws, row_num: int, index: int, user: User) -> None:
    ws.cell(row=row_num, column=1, value=index + 1)
    ws.cell(row=row_num, column=2, value=user.real_name if user.real_name is not None else "")
    ws.cell(row=row_num, column=3, value=_gender_to_string(user.gender))
    ws.cell(row=row_num, column=4, value=user.student_no if user.student_no is not None else "")
    ws.cell(row=row_num, column=5, value=user.school_name if user.school_name is not None else "")
    ws.cell(row=row_num, column=6, value=user.college_name if user.college_name is not None else "")
    ws.cell(row=row_num, column=7, value=user.class_name if user.class_name is not None else "")


def export_registration(db: Session, activity_id, current_user: User):
    activity = _export_context(db, activity_id, current_user)
    approved = _approved_registrations(db, activity_id)

    headers = ["序号", "姓名", "性别", "学号", "学校", "学院", "班级"]
    wb, ws = _new_workbook("报名情况表", headers)

    row_num = 2
    for index, reg in enumerate(approved):
        user = db.query(User).filter(User.id == reg.user_id).first()
        if user is None:
            continue
        _fill_user_columns(ws, row_num, index, user)
        row_num += 1

    filename = activity.name + "_报名情况表.xlsx"
    return _workbook_bytes(wb), filename


def export_checkin(db: Session, activity_id, current_user: User):
    activity = _export_context(db, activity_id, current_user)
    approved = _approved_registrations(db, activity_id)

    headers = ["序号", "姓名", "性别", "学号", "学校", "学院", "班级", "报名状态", "签到状态"]
    wb, ws = _new_workbook("签到情况表", headers)

    row_num = 2
    for index, reg in enumerate(approved):
        user = db.query(User).filter(User.id == reg.user_id).first()
        if user is None:
            continue
        checkin = (
            db.query(ActivityCheckin)
            .filter(
                ActivityCheckin.activity_id == activity_id,
                ActivityCheckin.user_id == reg.user_id,
            )
            .first()
        )
        _fill_user_columns(ws, row_num, index, user)
        ws.cell(row=row_num, column=8, value=reg.status if reg.status is not None else "")
        ws.cell(row=row_num, column=9, value="已签到" if checkin is not None else "未签到")
        row_num += 1

    filename = activity.name + "_签到情况表.xlsx"
    return _workbook_bytes(wb), filename


def export_score(db: Session, activity_id, current_user: User):
    activity = _export_context(db, activity_id, current_user)
    approved = _approved_registrations(db, activity_id)

    headers = ["序号", "姓名", "性别", "学号", "学校", "学院", "班级", "报名状态", "评分"]
    wb, ws = _new_workbook("评分情况表", headers)

    row_num = 2
    for index, reg in enumerate(approved):
        user = db.query(User).filter(User.id == reg.user_id).first()
        if user is None:
            continue
        _fill_user_columns(ws, row_num, index, user)
        ws.cell(row=row_num, column=8, value=reg.status if reg.status is not None else "")
        ws.cell(
            row=row_num,
            column=9,
            value=str(reg.score) if reg.score is not None else "未评分",
        )
        row_num += 1

    filename = activity.name + "_评分情况表.xlsx"
    return _workbook_bytes(wb), filename


def export_disposition(filename: str) -> str:
    return "attachment;filename=" + quote_plus(filename, safe="")


# ---------------------------------------------------------------------------
# 用户活动 / 二维码 / 签到 / 上下架 / 删除
# ---------------------------------------------------------------------------

def user_activities(
    db: Session, user_id, type_, keyword, current_user: User
):
    if user_id != current_user.id:
        assert_permission(db, current_user, Perm.USER_VIEW)

    if type_ is None or str(type_).strip() == "":
        type_ = "all"
    type_ = str(type_).lower().strip()

    if type_ not in ("participated", "managed", "all"):
        raise ValueError("type参数值无效，只能是：participated、managed 或 all")

    participated_dtos = []
    managed_dtos = []

    if type_ in ("participated", "all"):
        registrations = (
            db.query(ActivityRegistration)
            .filter(ActivityRegistration.user_id == user_id)
            .order_by(ActivityRegistration.registration_time.desc())
            .all()
        )

        registration_time_map = {}
        for reg in registrations:
            if reg.activity_id not in registration_time_map:
                registration_time_map[reg.activity_id] = reg.registration_time

        participated_ids = list({reg.activity_id for reg in registrations})

        participated_activities = []
        if participated_ids:
            query = db.query(Activity).filter(Activity.id.in_(participated_ids))
            if _has_text(keyword):
                kw = f"%{keyword}%"
                query = query.filter(
                    or_(Activity.name.like(kw), Activity.description.like(kw))
                )
            participated_activities = query.all()
            for activity in participated_activities:
                refresh_activity_status_by_time(db, activity)
                activity.current_participants = _approved_count(db, activity.id)

            def _compare_participated(a1, a2):
                t1 = registration_time_map.get(a1.id)
                t2 = registration_time_map.get(a2.id)
                if t1 is None and t2 is None:
                    return 0
                if t1 is None:
                    return 1
                if t2 is None:
                    return -1
                if t1 == t2:
                    return 0
                return -1 if t2 < t1 else 1

            participated_activities.sort(key=cmp_to_key(_compare_participated))

        participated_dtos = [serialize_activity_dto(a) for a in participated_activities]

    if type_ in ("managed", "all"):
        manager_records = (
            db.query(ClubMember)
            .filter(ClubMember.user_id == user_id, ClubMember.is_manager.is_(True))
            .all()
        )

        managed_activities = []
        if manager_records:
            managed_club_ids = list({m.club_id for m in manager_records})
            query = db.query(Activity).filter(Activity.club_id.in_(managed_club_ids))
            if _has_text(keyword):
                kw = f"%{keyword}%"
                query = query.filter(
                    or_(Activity.name.like(kw), Activity.description.like(kw))
                )
            query = query.order_by(Activity.create_time.desc())
            managed_activities = query.all()

            def _compare_managed(a1, a2):
                a1_running = a1.status == "进行中"
                a2_running = a2.status == "进行中"
                if a1_running and not a2_running:
                    return -1
                if not a1_running and a2_running:
                    return 1
                if a1.create_time == a2.create_time:
                    return 0
                return -1 if a2.create_time < a1.create_time else 1

            managed_activities.sort(key=cmp_to_key(_compare_managed))

            for activity in managed_activities:
                refresh_activity_status_by_time(db, activity)
                activity.current_participants = _approved_count(db, activity.id)

        managed_dtos = [serialize_activity_dto(a) for a in managed_activities]

    if type_ == "participated":
        data = {"activities": participated_dtos}
    elif type_ == "managed":
        data = {"activities": managed_dtos}
    else:
        data = {"participated": participated_dtos, "managed": managed_dtos}
    return success_response(data=data, message="查询成功")


def generate_qrcode(db: Session, activity_id, current_user: User):
    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ValueError("活动不存在")

    assert_permission(
        db,
        current_user,
        Perm.ACTIVITY_QRCODE,
        club_id=activity.club_id,
        message="无权限生成签到二维码",
    )

    now_ms = int(datetime.now().timestamp() * 1000)
    qr_content = json.dumps(
        {"activityId": activity_id, "timestamp": now_ms, "expiresIn": 60000}
    )

    qr = qrcode.QRCode(
        error_correction=qrcode.constants.ERROR_CORRECT_H, border=1
    )
    qr.add_data(qr_content)
    qr.make(fit=True)
    image = qr.make_image(fill_color="black", back_color="white").get_image()
    image = image.resize((300, 300))

    buffer = BytesIO()
    image.save(buffer, format="PNG")
    qr_base64 = "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()

    qrcode_data = {
        "qrCode": qr_base64,
        "activityId": activity_id,
        "activityName": activity.name,
        "expiresIn": 60000,
        "generateTime": int(datetime.now().timestamp() * 1000),
    }
    return success_response(data={"qrcode": qrcode_data}, message="生成二维码成功")


def checkin(db: Session, payload: dict, current_user: User):
    assert_permission(db, current_user, Perm.ACTIVITY_CHECKIN)

    user_id = current_user.id
    qr_content = payload.get("qrContent")
    if not _has_text(qr_content):
        raise ValueError("请扫描活动二维码签到")

    try:
        qr_data = json.loads(str(qr_content))
    except (TypeError, ValueError):
        raise ValueError("二维码内容格式不正确")

    activity_id = _to_long(qr_data.get("activityId"))
    if activity_id is None:
        raise ValueError("二维码中缺少活动ID")

    if "timestamp" in qr_data:
        timestamp = int(qr_data.get("timestamp"))
        expires_in = (
            int(qr_data.get("expiresIn")) if "expiresIn" in qr_data else 60000
        )
        now_ms = int(datetime.now().timestamp() * 1000)
        if now_ms - timestamp > expires_in:
            raise ValueError("二维码已过期，请刷新后重试")

    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ValueError("活动不存在")

    if activity.audit_status is None or activity.audit_status != 1:
        raise ValueError("活动尚未通过审核，无法签到")
    if activity.is_published is not True:
        raise ValueError("活动已下架，无法签到")

    registration = (
        db.query(ActivityRegistration)
        .filter(
            ActivityRegistration.activity_id == activity_id,
            ActivityRegistration.user_id == user_id,
        )
        .first()
    )
    if registration is None:
        raise ValueError("您未报名该活动，无法签到")
    if registration.status != "已通过":
        raise ValueError("您的报名尚未通过审核，无法签到")

    now = datetime.now()
    if activity.start_time is None or activity.end_time is None:
        raise ValueError("活动时间未配置，无法签到")
    if now < activity.start_time:
        raise ValueError("活动尚未开始，无法签到")
    if now > activity.end_time:
        raise ValueError("活动已结束，无法签到")

    existing_checkin = (
        db.query(ActivityCheckin)
        .filter(
            ActivityCheckin.activity_id == activity_id,
            ActivityCheckin.user_id == user_id,
        )
        .first()
    )
    if existing_checkin is not None:
        raise ValueError("您已经签到过了")

    checkin_row = ActivityCheckin()
    checkin_row.activity_id = activity_id
    checkin_row.user_id = user_id
    checkin_row.checkin_time = now
    db.add(checkin_row)
    db.commit()
    db.refresh(checkin_row)

    checkin_data = {
        "checkinId": checkin_row.id,
        "checkinTime": dt_to_millis(checkin_row.checkin_time),
        "activityName": activity.name,
    }
    return success_response(data={"checkin": checkin_data}, message="签到成功")


def update_publish_status(
    db: Session, activity_id, payload: dict, current_user: User
):
    if activity_id is None:
        raise ValueError("活动ID不能为空")

    is_published = None
    if "isPublished" in payload:
        is_published = str(payload.get("isPublished")).strip().lower() == "true"
    if is_published is None:
        raise ValueError("参数不完整，需要提供 isPublished")

    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ValueError("活动不存在")

    assert_permission(
        db,
        current_user,
        Perm.ACTIVITY_PUBLISH,
        club_id=activity.club_id,
        message="无权限操作该活动",
    )

    if is_published is True:
        if activity.audit_status is None or activity.audit_status != 1:
            raise ValueError("活动尚未通过审核，无法上架")

    activity.is_published = is_published
    db.commit()
    db.refresh(activity)

    message = "活动上架成功" if is_published is True else "活动下架成功"
    return success_response(
        data={"activity": serialize_activity(activity)}, message=message
    )


def delete_activity(db: Session, activity_id, current_user: User):
    if activity_id is None:
        raise ValueError("活动ID不能为空")

    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if activity is None:
        raise ValueError("活动不存在")

    assert_permission(
        db,
        current_user,
        Perm.ACTIVITY_DELETE,
        club_id=activity.club_id,
        message="无权限删除该活动",
    )

    db.delete(activity)
    db.commit()
    return success_response(message="删除成功")
